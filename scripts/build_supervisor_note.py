"""Generate a supervisor correction note from saved measurements."""

from __future__ import annotations

import argparse
from datetime import date
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from scripts.corrected_campaign import load_corrected_a4, validate_campaign


def build_note(sac, campaign=None, a4=None, robustness=None, board=None,
               accel=None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.75)
    doc.add_paragraph("Correction to the EV energy management results", "Title")
    doc.add_paragraph(date.today().strftime("%d %B %Y").lstrip("0"))
    doc.add_paragraph(
        "The earlier claim of an 8.28% mean energy saving is withdrawn. "
        "The host vehicle model capped the energy billed at the commanded "
        "power ceiling while delivering unrestricted torque. The corrected "
        "model reduces torque when the power ceiling binds. Regression tests "
        "cover this correction.")
    doc.add_heading("Retraining on the corrected model", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for c, text in zip(table.rows[0].cells,
                       ("Controller", "Validation energy Wh/km", "Against rules")):
        c.text = text
    rows = [("Rule manager", f"{sac['baseline']['equivalent_wh_per_km']:.2f}", "")]
    for arm in ("without-forecast", "with-forecast"):
        s = sac["summary"][arm]
        rows.append((arm.replace("-", " "), f"{s['mean']:.2f} ± {s['std']:.2f}",
                     f"{s['gap_to_rule_pct_mean']:+.2f}%"))
    for row in rows:
        for c, text in zip(table.add_row().cells, row):
            c.text = text
    doc.add_paragraph(
        f"SIMULATED on {sac['cycle'].upper()}, {sac['seeds']} seeds per "
        "learned arm; spread is the seed standard deviation. Shed cabin heat "
        "is charged back. The no-forecast policy was selected on validation "
        "and exported. The paired forecast interval "
        + ("includes zero." if sac["a3"]["null"] else "excludes zero."))
    doc.add_heading("Board measurements", level=1)
    if campaign is None:
        doc.add_paragraph(
            "The corrected board campaign and link-outage sweep are pending. "
            "The simulation results above are not board measurements. This "
            "note will be regenerated when the saved measurements are complete.")
    else:
        summary = validate_campaign(campaign)
        doc.add_paragraph(
            f"MEASURED: {summary['runs']} runs, {summary['steps']:,} control "
            f"steps, {summary['ctrl_late']} late steps. The learned controller "
            f"uses less energy in {summary['actor_less']} of "
            f"{summary['valid_settings']} valid settings, more in "
            f"{summary['actor_more']}, with {summary['ties']} ties. The mean "
            f"paired difference is {summary['mean_gap_pct']:+.3f}%. Each cell "
            "has one run per controller, so these settings are not repeats.")
        doc.add_paragraph(
            f"{summary['excluded_settings']} settings are excluded from energy "
            "aggregates because tracking RMSE reaches 1 m/s; their rows remain "
            "visible. The serialized delay axis measures host throughput, "
            "not independent link latency.")
        if a4 is None:
            doc.add_paragraph("The corrected link-outage sweep remains pending.")
        else:
            late = sum(d["delta"].get("ctrl_late", 0)
                       for pair in a4.values() for d in pair)
            failed = sum(d.get("verdict") is not True
                         for pair in a4.values() for d in pair)
            doc.add_paragraph(
                f"Corrected A4: {len(a4)} outage lengths from 0.2 to 10 s, "
                f"eight paired onsets each, {late} late steps and {failed} "
                "files with counter checks not passed. The report presents "
                "the new sweep separately from historical measurements.")
    doc.add_heading("Speed-loop robustness", level=1)
    doc.add_paragraph(
        "The proposed speed-loop decision transformer was dropped: in "
        "simulation a Kalman filter on measured speed met the criteria set "
        "for it. The filter "
        "and a speed-sensor fault detector are in the firmware, switched off "
        "by default.")
    if board:
        rows = {(r["scenario"], r["level"]): r for r in board["rows"]}
        fl, sv = rows[("sensor_noise", "floor")], rows[("sensor_noise", "severe")]
        v5, v6 = board["verdicts"]["b5"], board["verdicts"]["b6"]
        doc.add_paragraph(
            "MEASURED on the board, rule manager for energy in every arm, one "
            "run per cell. Ten times the sensor noise costs the shipped speed "
            f"loop {v5['b0 penalty at 10x noise, Wh/km']:.2f} Wh/km on US06 "
            f"({fl['b0_equivalent_wh_per_km']:.2f} to "
            f"{sv['b0_equivalent_wh_per_km']:.2f}). With the filter it costs "
            f"{v5['own penalty at 10x noise, Wh/km']:.2f}, and with the "
            f"detector added {v6['own penalty at 10x noise, Wh/km']:.2f}. "
            f"Peak battery current under that noise falls from "
            f"{sv['b0_peak_current_a']:.0f} A to "
            f"{sv['b5_peak_current_a']:.0f} A. Worst control computation "
            f"rises from {board['b0 worst control computation, us']} to "
            f"{v6['worst control computation, us']} microseconds of a 10,000 "
            "microsecond period, with no late step.")
        doc.add_paragraph(
            "Against the criteria fixed beforehand, each arm misses one. The "
            "filter misses the clean-tracking bound on the urban cycle by "
            f"{v5['clean tracking RMSE minus b0, m/s, per cycle'][0] - 0.005:.3f} "
            "m/s, which is less than the shipped loop differs from itself "
            "between two identical runs. The detector raised its flag "
            f"{v6['false alarms on a healthy sensor']} time on a healthy "
            f"sensor in {v6['hours on the board']:.2f} hours. The alarm "
            "recurs on the same control step in two further board runs and "
            "in a replay on the laptop, so it is not a link or chip fault: "
            "the detector's threshold allows for noise on the speed reading "
            "but not for drift in the speed it works out from acceleration.")
    if robustness:
        v = robustness["verdict"]
        doc.add_paragraph(
            f"SIMULATED, detector: {v['named faults run']} named faults, "
            f"{v['named faults missed']} missed, worst delay "
            f"{v['worst delay, s']:.2f} s; "
            f"{v['false alarms in them']} false alarms in "
            f"{v['clean hours at the noise floor']:.2f} h at the noise floor. "
            "No fault was injected on the board, so fault detection is not a "
            "board result. A 0.5 m/s offset under ten times the noise is "
            "flagged but not corrected.")
    if accel:
        s = accel["summary"]
        doc.add_paragraph(
            "The detector does not survive a real acceleration channel as "
            f"tuned. On {s['trips']} logged BMW i3 trips ({s['hours']:.0f} "
            "hours, no sensor fault in the data) logged acceleration and "
            "speed disagree by a slow component of about "
            f"{s['slow_disagreement_sd_mps2_median']:.2f} m/s^2, which the "
            "detector's dead reckoning turns into a standing gap several "
            "times its fault threshold: it would be flagged "
            f"{100 * s['flagged_share']:.0f}% of the time. The simulated "
            "channel carried white noise only and did not show this. The "
            "filter does not use the acceleration channel and is unaffected.")
    doc.add_paragraph("Draft for review; this script does not send the note.")
    return doc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="reports/supervisor-correction-note.docx")
    ap.add_argument("--robustness-docs", type=Path)
    args = ap.parse_args()
    docs = Path("docs")
    sac = json.loads((docs / "week11-sac-corrected.json").read_text())
    campaign_path = docs / "week11-pil-corrected.json"
    campaign = json.loads(campaign_path.read_text()) if campaign_path.exists() else None
    robustness = board = accel = None
    if args.robustness_docs:
        # The detector was corrected on 2026-10-07; its evaluation after the
        # correction is the week-12 file, and the earlier one is kept.
        newer = args.robustness_docs / "week12-fault-arms-b6.json"
        older = args.robustness_docs / "week11-fault-arms-b6.json"
        robustness = json.loads((newer if newer.exists() else older).read_text())
        board_path = args.robustness_docs / "week12-speed-arms.json"
        if board_path.exists():
            board = json.loads(board_path.read_text())
        accel_path = args.robustness_docs / "week12-accel-channel-check.json"
        if accel_path.exists():
            accel = json.loads(accel_path.read_text())
    doc = build_note(sac, campaign, load_corrected_a4(docs), robustness, board,
                     accel)
    path = Path(args.out)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
