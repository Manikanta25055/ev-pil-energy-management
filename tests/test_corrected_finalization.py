"""Checks that the corrected campaign cannot be finalized from stale data."""

from __future__ import annotations

import copy
import importlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from evpil.control.scenarios import AXES


class CorrectedFinalizationTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec("scripts.corrected_campaign"),
                             "corrected campaign validation is missing")
        self.module = importlib.import_module("scripts.corrected_campaign")
        self.campaign = {"runs": []}
        for controller in ("rule", "actor"):
            for axis, levels in AXES.items():
                for seed, lv in enumerate(levels):
                    self.campaign["runs"].append({
                        "scenario": axis, "level": lv.level,
                        "controller": controller, "seed": seed,
                        "cycle": lv.cycle, "grade_offset_rad": lv.grade_offset_rad,
                        "soc0": lv.soc0, "noise_scale": lv.noise_scale,
                        "comm_delay_s": lv.comm_delay_s, "aggression": lv.aggression,
                        "steps": 100, "duration_s": 1.0, "distance_km": 0.02,
                        "equivalent_wh_per_km": 190.0, "tracking_rmse_mps": 0.02,
                        "ctrl_late": 0, "dropped_frames": 0,
                        "actor_steps": 10 if controller == "actor" else 0,
                    })

    def test_complete_campaign_accepts_invalid_delay_cells_as_visible_measurements(self):
        for r in self.campaign["runs"]:
            if r["comm_delay_s"] > 0:
                r.update(tracking_rmse_mps=30.0, dropped_frames=14000)
        summary = self.module.validate_campaign(self.campaign)
        self.assertEqual(summary["valid_settings"], 16)
        self.assertEqual(summary["excluded_settings"], 2)

    def test_duplicate_cell_cannot_replace_a_missing_cell(self):
        self.campaign["runs"][-1] = copy.deepcopy(self.campaign["runs"][0])
        with self.assertRaisesRegex(ValueError, "duplicate|cells"):
            self.module.validate_campaign(self.campaign)

    def test_wrong_firmware_arm_is_rejected(self):
        self.campaign["runs"][0]["actor_steps"] = 1
        with self.assertRaisesRegex(ValueError, "firmware|actor_steps"):
            self.module.validate_campaign(self.campaign)

    def test_wrong_scenario_parameters_are_rejected(self):
        self.campaign["runs"][0]["noise_scale"] = 25.0
        with self.assertRaisesRegex(ValueError, "noise_scale"):
            self.module.validate_campaign(self.campaign)

    def test_disturbed_valid_cell_is_not_finalized(self):
        self.campaign["runs"][0]["dropped_frames"] = 54
        with self.assertRaisesRegex(ValueError, "disturbed"):
            self.module.validate_campaign(self.campaign)

    def test_nonfinite_energy_is_rejected(self):
        self.campaign["runs"][0]["equivalent_wh_per_km"] = float("nan")
        with self.assertRaisesRegex(ValueError, "finite"):
            self.module.validate_campaign(self.campaign)

    def test_ties_are_not_counted_as_energy_wins(self):
        summary = self.module.validate_campaign(self.campaign)
        self.assertEqual(summary["actor_less"], 0)
        self.assertEqual(summary["actor_more"], 0)
        self.assertEqual(summary["ties"], 18)

    def test_corrected_a4_never_falls_back_to_week9_measurements(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td)
            source = Path("docs/week9-a4-arduino-on-0.2s.json").read_bytes()
            (p / "week9-a4-arduino-on-0.2s.json").write_bytes(source)
            self.assertIsNone(self.module.load_corrected_a4(p))

    def test_complete_corrected_a4_requires_identical_onsets(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td)
            for mode in ("on", "off"):
                for s in self.module.DURATIONS:
                    d = json.loads(Path(f"docs/week9-a4-arduino-{mode}-{s}s.json").read_text())
                    (p / f"week11-a4-{mode}-{s}s.json").write_text(json.dumps(d))
            self.assertEqual(len(self.module.load_corrected_a4(p)), 8)
            bad = p / "week11-a4-off-0.2s.json"
            d = json.loads(bad.read_text())
            d["blackouts"][0]["onset_s"] += 1
            bad.write_text(json.dumps(d))
            with self.assertRaisesRegex(ValueError, "onsets"):
                self.module.load_corrected_a4(p)


class CorrectedReportTests(unittest.TestCase):
    def test_rendered_corrected_campaign_table_matches_saved_measurements(self):
        campaign_path = Path("docs/week11-pil-corrected.json")
        if not campaign_path.exists():
            self.skipTest("corrected board campaign is still running")
        from docx import Document
        data = json.loads(campaign_path.read_text())
        report = Path("reports/EV-PIL-status-report.docx")
        self.assertGreaterEqual(report.stat().st_mtime, campaign_path.stat().st_mtime)
        doc = Document(report)
        header = ["Varied", "Setting", "Rules: tracking error, m/s",
                  "Learned: tracking error, m/s", "Rules: energy, Wh/km",
                  "Learned: energy, Wh/km", "Learned vs rules"]
        tables = [t for t in doc.tables if [c.text for c in t.rows[0].cells] == header]
        self.assertEqual(len(tables), 1)
        by_level = {(r["level"], r["controller"]): r for r in data["runs"]}
        self.assertEqual(len(tables[0].rows) - 1, 18)
        for row in tables[0].rows[1:]:
            cells = [c.text for c in row.cells]
            lv = cells[1].replace(" ", "_")
            rule, actor = (by_level[(lv, arm)] for arm in ("rule", "actor"))
            self.assertEqual(cells[2:6], [f"{rule['tracking_rmse_mps']:.4f}",
                                        f"{actor['tracking_rmse_mps']:.4f}",
                                        f"{rule['equivalent_wh_per_km']:.2f}",
                                        f"{actor['equivalent_wh_per_km']:.2f}"])

    def test_supervisor_note_cannot_present_pending_board_results_as_measured(self):
        self.assertIsNotNone(importlib.util.find_spec("scripts.build_supervisor_note"))
        from scripts import build_supervisor_note
        self.assertTrue(hasattr(build_supervisor_note, "build_note"),
                        "supervisor correction note is not generated from data")
        data = json.loads(Path("docs/week11-sac-corrected.json").read_text())
        doc = build_supervisor_note.build_note(data)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("pending", text)
        self.assertIn("withdrawn", text)
        self.assertIn("SIMULATED", text)
        self.assertNotIn("36 measured runs", text)

    def test_report_contains_retraining_numbers_from_the_new_file(self):
        from docx import Document
        from scripts import build_report
        self.assertTrue(hasattr(build_report, "retraining_section"),
                        "report still omits corrected-plant retraining")
        data = json.loads(Path("docs/week11-sac-corrected.json").read_text())
        doc = Document()
        build_report.retraining_section(doc, data)
        text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn(f"{data['baseline']['equivalent_wh_per_km']:.2f}", text)
        for arm in ("with-forecast", "without-forecast"):
            self.assertIn(f"{data['summary'][arm]['mean']:.2f}", text)
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("SIMULATED", paragraphs)
        self.assertIn("validation", paragraphs)


if __name__ == "__main__":
    unittest.main()
